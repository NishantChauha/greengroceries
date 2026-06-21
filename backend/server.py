from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import logging
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Annotated
from io import StringIO
import csv

import bcrypt
import jwt
from bson import ObjectId
from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr, BeforeValidator, ConfigDict

# ----- MongoDB -----
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_ALGORITHM = "HS256"
JWT_SECRET = os.environ["JWT_SECRET"]

# ----- Helpers -----
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

def create_access_token(user_id: str, email: str, role: str) -> str:
    payload = {
        "sub": user_id, "email": email, "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(hours=24),
        "type": "access",
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


def PyObjectId():
    return Annotated[str, BeforeValidator(lambda v: str(v) if v else v)]


# ----- Models -----
class RegisterIn(BaseModel):
    email: EmailStr
    password: str = Field(min_length=6)
    name: str = Field(min_length=2)
    hotel_name: str = Field(min_length=2)
    phone: Optional[str] = None
    address: Optional[str] = None

class LoginIn(BaseModel):
    email: EmailStr
    password: str

class UserOut(BaseModel):
    id: str
    email: str
    name: str
    role: str
    hotel_name: Optional[str] = None
    phone: Optional[str] = None
    address: Optional[str] = None

class ItemIn(BaseModel):
    name: str = Field(min_length=1)
    category: str = Field(default="vegetable")  # vegetable | fruit
    unit: str = Field(default="kg")
    icon: Optional[str] = None

class ItemOut(BaseModel):
    id: str
    name: str
    category: str
    unit: str
    icon: Optional[str] = None

class OrderLine(BaseModel):
    item_id: str
    name: str
    unit: str
    category: str
    quantity: float
    rate: float = 0.0          # price per unit (admin-entered)
    amount: float = 0.0        # quantity * rate (computed)

class OrderIn(BaseModel):
    order_date: str  # YYYY-MM-DD
    notes: Optional[str] = None
    lines: List[OrderLine]


class OrderUpdateIn(BaseModel):
    lines: Optional[List[OrderLine]] = None
    notes: Optional[str] = None
    tax_rate: Optional[float] = None  # percent, e.g. 5


def _compute_totals(lines: List[dict], tax_rate: float):
    subtotal = 0.0
    for ln in lines:
        amount = float(ln.get("quantity", 0)) * float(ln.get("rate", 0))
        ln["amount"] = round(amount, 2)
        subtotal += amount
    subtotal = round(subtotal, 2)
    tax = round(subtotal * float(tax_rate) / 100.0, 2)
    grand_total = round(subtotal + tax, 2)
    return subtotal, tax, grand_total

class OrderOut(BaseModel):
    id: str
    hotel_id: str
    hotel_name: str
    order_date: str
    notes: Optional[str] = None
    status: str
    created_at: str
    lines: List[OrderLine]
    tax_rate: float = 0.0
    subtotal: float = 0.0
    tax: float = 0.0
    grand_total: float = 0.0


# ----- App -----
app = FastAPI(title="Green Groceries API")
api = APIRouter(prefix="/api")


async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Invalid token type")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="User not found")
        user["id"] = str(user["_id"])
        user.pop("_id", None)
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")


def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin only")
    return user


def set_auth_cookies(response: Response, token: str):
    response.set_cookie(
        key="access_token", value=token, httponly=True,
        secure=False, samesite="lax", max_age=86400, path="/",
    )


def user_to_out(user: dict) -> dict:
    return {
        "id": user.get("id") or str(user.get("_id")),
        "email": user["email"],
        "name": user["name"],
        "role": user["role"],
        "hotel_name": user.get("hotel_name"),
        "phone": user.get("phone"),
        "address": user.get("address"),
    }


# ----- Auth endpoints -----
# Public registration is disabled — hotels are created by admin.

@api.post("/auth/login")
async def login(payload: LoginIn, response: Response):
    email = payload.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    user_id = str(user["_id"])
    token = create_access_token(user_id, email, user["role"])
    set_auth_cookies(response, token)
    user["id"] = user_id
    return {"user": user_to_out(user), "access_token": token}


@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api.get("/auth/me", response_model=UserOut)
async def me(user: dict = Depends(get_current_user)):
    return user_to_out(user)


# ----- Items -----
@api.get("/items", response_model=List[ItemOut])
async def list_items(user: dict = Depends(get_current_user)):
    docs = await db.items.find({}).sort("name", 1).to_list(1000)
    return [{"id": str(d["_id"]), "name": d["name"], "category": d["category"],
             "unit": d["unit"], "icon": d.get("icon")} for d in docs]


@api.post("/items", response_model=ItemOut)
async def create_item(payload: ItemIn, _: dict = Depends(require_admin)):
    existing = await db.items.find_one({"name": payload.name})
    if existing:
        raise HTTPException(status_code=400, detail="Item with this name already exists")
    doc = payload.model_dump()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    result = await db.items.insert_one(doc)
    return {"id": str(result.inserted_id), **payload.model_dump()}


@api.put("/items/{item_id}", response_model=ItemOut)
async def update_item(item_id: str, payload: ItemIn, _: dict = Depends(require_admin)):
    await db.items.update_one({"_id": ObjectId(item_id)}, {"$set": payload.model_dump()})
    doc = await db.items.find_one({"_id": ObjectId(item_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Item not found")
    return {"id": str(doc["_id"]), "name": doc["name"], "category": doc["category"],
            "unit": doc["unit"], "icon": doc.get("icon")}


@api.delete("/items/{item_id}")
async def delete_item(item_id: str, _: dict = Depends(require_admin)):
    result = await db.items.delete_one({"_id": ObjectId(item_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Item not found")
    return {"ok": True}


# ----- Orders -----
@api.post("/orders", response_model=OrderOut)
async def create_order(payload: OrderIn, user: dict = Depends(get_current_user)):
    if user["role"] != "hotel":
        raise HTTPException(status_code=403, detail="Only hotels can place orders")
    if not payload.lines:
        raise HTTPException(status_code=400, detail="At least one item required")
    valid_lines = [ln.model_dump() for ln in payload.lines if ln.quantity > 0]
    if not valid_lines:
        raise HTTPException(status_code=400, detail="Add quantity for at least one item")
    subtotal, tax, grand = _compute_totals(valid_lines, 0.0)
    doc = {
        "hotel_id": user["id"],
        "hotel_name": user.get("hotel_name") or user["name"],
        "order_date": payload.order_date,
        "notes": payload.notes,
        "status": "pending",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "lines": valid_lines,
        "tax_rate": 0.0,
        "subtotal": subtotal, "tax": tax, "grand_total": grand,
    }
    result = await db.orders.insert_one(doc)
    doc["id"] = str(result.inserted_id)
    doc.pop("_id", None)
    return doc


def _serialize_order(d: dict) -> dict:
    return {
        "id": str(d["_id"]), "hotel_id": d["hotel_id"], "hotel_name": d["hotel_name"],
        "order_date": d["order_date"], "notes": d.get("notes"), "status": d["status"],
        "created_at": d["created_at"], "lines": d["lines"],
        "tax_rate": float(d.get("tax_rate") or 0.0),
        "subtotal": float(d.get("subtotal") or 0.0),
        "tax": float(d.get("tax") or 0.0),
        "grand_total": float(d.get("grand_total") or 0.0),
    }


@api.get("/orders", response_model=List[OrderOut])
async def list_orders(
    user: dict = Depends(get_current_user),
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    hotel_id: Optional[str] = None,
):
    query: dict = {}
    if user["role"] == "hotel":
        query["hotel_id"] = user["id"]
    elif hotel_id:
        query["hotel_id"] = hotel_id
    if date_from and date_to:
        query["order_date"] = {"$gte": date_from, "$lte": date_to}
    elif date_from:
        query["order_date"] = {"$gte": date_from}
    elif date_to:
        query["order_date"] = {"$lte": date_to}
    if status:
        query["status"] = status
    docs = await db.orders.find(query).sort([("order_date", -1), ("created_at", -1)]).to_list(1000)
    return [_serialize_order(d) for d in docs]


@api.put("/orders/{order_id}", response_model=OrderOut)
async def update_order(order_id: str, payload: OrderUpdateIn, _: dict = Depends(require_admin)):
    existing = await db.orders.find_one({"_id": ObjectId(order_id)})
    if not existing:
        raise HTTPException(status_code=404, detail="Order not found")

    new_lines = ([ln.model_dump() for ln in payload.lines] if payload.lines is not None
                 else existing["lines"])
    tax_rate = float(payload.tax_rate if payload.tax_rate is not None else existing.get("tax_rate") or 0.0)
    subtotal, tax, grand = _compute_totals(new_lines, tax_rate)
    updates = {
        "lines": new_lines, "tax_rate": tax_rate,
        "subtotal": subtotal, "tax": tax, "grand_total": grand,
    }
    if payload.notes is not None:
        updates["notes"] = payload.notes
    await db.orders.update_one({"_id": ObjectId(order_id)}, {"$set": updates})
    doc = await db.orders.find_one({"_id": ObjectId(order_id)})
    return _serialize_order(doc)


@api.patch("/orders/{order_id}/status")
async def update_order_status(order_id: str, body: dict, user: dict = Depends(get_current_user)):
    new_status = body.get("status")
    if new_status not in {"pending", "delivered", "cancelled"}:
        raise HTTPException(status_code=400, detail="Invalid status")

    # Hotels may only cancel their own pending orders. Admin may set any status.
    if user["role"] == "hotel":
        if new_status != "cancelled":
            raise HTTPException(status_code=403, detail="Hotels can only cancel orders")
        existing = await db.orders.find_one({"_id": ObjectId(order_id)})
        if not existing:
            raise HTTPException(status_code=404, detail="Order not found")
        if existing["hotel_id"] != user["id"]:
            raise HTTPException(status_code=403, detail="Not your order")
        if existing["status"] != "pending":
            raise HTTPException(status_code=400, detail="Only pending orders can be cancelled")

    result = await db.orders.update_one(
        {"_id": ObjectId(order_id)}, {"$set": {"status": new_status}}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Order not found")
    return {"ok": True, "status": new_status}


@api.delete("/orders/{order_id}")
async def delete_order(order_id: str, _: dict = Depends(require_admin)):
    # Hard delete is admin-only. Hotels should use PATCH /status cancelled to preserve history.
    result = await db.orders.delete_one({"_id": ObjectId(order_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Order not found")
    return {"ok": True}


# ----- Purchase Sheet -----
@api.get("/purchase-sheet")
async def purchase_sheet(
    date_from: str,
    date_to: Optional[str] = None,
    _: dict = Depends(require_admin),
):
    if not date_to:
        date_to = date_from
    docs = await db.orders.find(
        {"order_date": {"$gte": date_from, "$lte": date_to}, "status": {"$ne": "cancelled"}}
    ).to_list(5000)

    agg: dict = {}
    hotels: set = set()
    total_orders = 0
    for d in docs:
        total_orders += 1
        hotels.add(d["hotel_id"])
        for ln in d["lines"]:
            key = ln["name"]
            if key not in agg:
                agg[key] = {
                    "name": ln["name"], "category": ln["category"],
                    "unit": ln["unit"], "total_quantity": 0.0, "hotel_count": set(),
                }
            agg[key]["total_quantity"] += float(ln["quantity"])
            agg[key]["hotel_count"].add(d["hotel_id"])

    rows = []
    for v in agg.values():
        rows.append({
            "name": v["name"], "category": v["category"], "unit": v["unit"],
            "total_quantity": round(v["total_quantity"], 2),
            "hotel_count": len(v["hotel_count"]),
        })
    rows.sort(key=lambda r: (r["category"], r["name"]))

    return {
        "date_from": date_from,
        "date_to": date_to,
        "total_orders": total_orders,
        "total_hotels": len(hotels),
        "rows": rows,
    }


@api.get("/purchase-sheet/export")
async def export_purchase_sheet(
    date_from: str,
    date_to: Optional[str] = None,
    fmt: str = "csv",
    _: dict = Depends(require_admin),
):
    sheet = await purchase_sheet(date_from=date_from, date_to=date_to, _=_)
    if fmt == "pdf":
        return _pdf_purchase_sheet(sheet)
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Green Groceries - Combined Purchase Sheet"])
    writer.writerow([f"From: {sheet['date_from']}", f"To: {sheet['date_to']}"])
    writer.writerow([f"Hotels: {sheet['total_hotels']}", f"Orders: {sheet['total_orders']}"])
    writer.writerow([])
    writer.writerow(["Item", "Category", "Unit", "Total Quantity", "Hotels Ordered"])
    for r in sheet["rows"]:
        writer.writerow([r["name"], r["category"], r["unit"], r["total_quantity"], r["hotel_count"]])
    buf.seek(0)
    filename = f"purchase_sheet_{sheet['date_from']}_to_{sheet['date_to']}.csv"
    return StreamingResponse(
        iter([buf.read()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ----- Orders export (CSV / PDF) — supports per-hotel filter -----
async def _gather_orders(user: dict, date_from: Optional[str], date_to: Optional[str],
                        status: Optional[str], hotel_id: Optional[str]) -> List[dict]:
    query: dict = {}
    if user["role"] == "hotel":
        query["hotel_id"] = user["id"]
    elif hotel_id:
        query["hotel_id"] = hotel_id
    if date_from and date_to:
        query["order_date"] = {"$gte": date_from, "$lte": date_to}
    elif date_from:
        query["order_date"] = {"$gte": date_from}
    elif date_to:
        query["order_date"] = {"$lte": date_to}
    if status:
        query["status"] = status
    return await db.orders.find(query).sort([("order_date", -1), ("created_at", -1)]).to_list(5000)


@api.get("/orders/export")
async def export_orders(
    fmt: str = "csv",
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    hotel_id: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    docs = await _gather_orders(user, date_from, date_to, status, hotel_id)
    label = "all"
    if hotel_id and user["role"] == "admin":
        h = await db.users.find_one({"_id": ObjectId(hotel_id)})
        if h:
            label = (h.get("hotel_name") or "hotel").replace(" ", "_")
    elif user["role"] == "hotel":
        label = (user.get("hotel_name") or "myhotel").replace(" ", "_")

    range_label = f"{date_from or 'all'}_to_{date_to or 'all'}"
    if fmt == "pdf":
        return _pdf_orders(docs, label, date_from, date_to)
    if fmt == "xlsx":
        return _xlsx_orders(docs, label, date_from, date_to)

    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Green Groceries - Orders Export"])
    writer.writerow([f"Hotel: {label}", f"From: {date_from or '—'}", f"To: {date_to or '—'}", f"Status: {status or 'all'}"])
    writer.writerow([])
    writer.writerow(["Order Date", "Hotel", "Item", "Category", "Unit", "Quantity", "Rate", "Amount", "Status", "Notes", "Placed At"])
    for d in docs:
        for ln in d["lines"]:
            writer.writerow([
                d["order_date"], d["hotel_name"], ln["name"], ln["category"], ln["unit"],
                ln.get("quantity", 0), ln.get("rate", 0), ln.get("amount", 0),
                d["status"], d.get("notes") or "", d["created_at"],
            ])
    buf.seek(0)
    filename = f"orders_{label}_{range_label}.csv"
    return StreamingResponse(
        iter([buf.read()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _xlsx_orders(docs: List[dict], label: str, date_from, date_to):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"

    head_fill = PatternFill("solid", fgColor="1F4A2C")
    head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4A2C")
    sub_font = Font(name="Calibri", size=10, color="747A76")
    money_fmt = "#,##0.00"
    thin = Side(border_style="thin", color="E5E0D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Green Groceries — Orders Report"
    ws["A1"].font = title_font
    ws.merge_cells("A1:I1")
    ws["A2"] = f"Hotel: {label}    From: {date_from or '—'}    To: {date_to or '—'}"
    ws["A2"].font = sub_font
    ws.merge_cells("A2:I2")

    headers = ["Order Date", "Hotel", "Item", "Unit", "Quantity", "Rate", "Amount", "Status", "Notes"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = border

    row = 5
    grand_total = 0.0
    for d in docs:
        for ln in d["lines"]:
            qty = float(ln.get("quantity", 0))
            rate = float(ln.get("rate", 0))
            amt = float(ln.get("amount", qty * rate))
            grand_total += amt
            values = [d["order_date"], d["hotel_name"], ln["name"], ln["unit"],
                      qty, rate, amt, d["status"], d.get("notes") or ""]
            for col, v in enumerate(values, 1):
                c = ws.cell(row=row, column=col, value=v)
                c.border = border
                if col in (5, 6, 7):
                    c.number_format = money_fmt
                    c.alignment = Alignment(horizontal="right")
            row += 1

    # Totals row
    tot_row = row + 1
    ws.cell(row=tot_row, column=6, value="Grand Total").font = Font(bold=True)
    c = ws.cell(row=tot_row, column=7, value=round(grand_total, 2))
    c.font = Font(bold=True, color="1F4A2C")
    c.number_format = money_fmt

    widths = [13, 22, 18, 8, 11, 11, 14, 12, 30]
    for i, w in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"orders_{label}.xlsx"
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ----- Invoice PDF -----
@api.get("/orders/{order_id}/invoice.pdf")
async def order_invoice(order_id: str, user: dict = Depends(get_current_user)):
    doc = await db.orders.find_one({"_id": ObjectId(order_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Order not found")
    if user["role"] == "hotel" and doc["hotel_id"] != user["id"]:
        raise HTTPException(status_code=403, detail="Not your order")
    hotel = await db.users.find_one({"_id": ObjectId(doc["hotel_id"])}) if doc.get("hotel_id") else None
    return _pdf_invoice(doc, hotel or {})


def _pdf_invoice(order: dict, hotel: dict):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from io import BytesIO

    out = BytesIO()
    doc_pdf = SimpleDocTemplate(out, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm,
                                topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    brand = ParagraphStyle("brand", parent=styles["Heading1"], fontSize=22,
                           textColor=colors.HexColor("#1F4A2C"), leading=24)
    sub = ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.HexColor("#747A76"),
                         fontSize=9, leading=12)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12,
                        textColor=colors.HexColor("#1F4A2C"))
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=10, leading=13)

    inv_no = f"GG-{order.get('order_date','').replace('-','')}-{str(order.get('_id') or order.get('id'))[-6:]}"
    issued = datetime.now(timezone.utc).strftime("%d %b %Y")

    elems = []
    # Header band
    header_tbl = Table([
        [
            Paragraph("Green Groceries", brand),
            Paragraph(
                f"<b>INVOICE</b><br/>"
                f"Invoice No: <b>{inv_no}</b><br/>"
                f"Order Date: <b>{order.get('order_date','—')}</b><br/>"
                f"Issued: {issued}", body),
        ]
    ], colWidths=[100*mm, 70*mm])
    header_tbl.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
    ]))
    elems.append(header_tbl)
    elems.append(Spacer(1, 4*mm))
    elems.append(Paragraph("Hotel produce ordering · Daily mandi sheet", sub))
    elems.append(Spacer(1, 8*mm))

    # Bill to / from
    bill_to = (
        f"<b>{order.get('hotel_name','—')}</b><br/>"
        f"{hotel.get('name','') or ''}<br/>"
        f"{hotel.get('email','') or ''}<br/>"
        f"{hotel.get('phone','') or ''}<br/>"
        f"{hotel.get('address','') or ''}"
    )
    bill_block = Table([
        [Paragraph("BILL TO", h2), Paragraph("FROM", h2)],
        [Paragraph(bill_to, body),
         Paragraph("Green Groceries<br/>Wholesale Fruits & Vegetables<br/>orders@greengroceries.com", body)],
    ], colWidths=[85*mm, 85*mm])
    bill_block.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    elems.append(bill_block)
    elems.append(Spacer(1, 8*mm))

    # Items table
    data = [["#", "Product", "Unit", "Qty", "Rate", "Amount"]]
    subtotal = 0.0
    for i, ln in enumerate(order["lines"], 1):
        qty = float(ln.get("quantity", 0))
        rate = float(ln.get("rate", 0))
        amount = float(ln.get("amount", qty * rate))
        subtotal += amount
        data.append([
            str(i), ln["name"], ln.get("unit", ""),
            f"{qty:g}", f"{rate:,.2f}", f"{amount:,.2f}",
        ])

    t = Table(data, repeatRows=1, hAlign="LEFT",
              colWidths=[10*mm, 70*mm, 18*mm, 22*mm, 25*mm, 30*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4A2C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.HexColor("#FFFFFF"), colors.HexColor("#F7F5F0")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E0D8")),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elems.append(t)
    elems.append(Spacer(1, 4*mm))

    # Totals box
    tax_rate = float(order.get("tax_rate") or 0)
    tax = float(order.get("tax") or round(subtotal * tax_rate / 100.0, 2))
    grand = float(order.get("grand_total") or round(subtotal + tax, 2))

    totals_data = [
        ["Subtotal", f"₹ {subtotal:,.2f}"],
        [f"Tax ({tax_rate:g}%)", f"₹ {tax:,.2f}"],
        ["Grand Total", f"₹ {grand:,.2f}"],
    ]
    totals = Table(totals_data, colWidths=[40*mm, 40*mm], hAlign="RIGHT")
    totals.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TEXTCOLOR", (0, 2), (-1, 2), colors.HexColor("#1F4A2C")),
        ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
        ("LINEABOVE", (0, 2), (-1, 2), 0.8, colors.HexColor("#1F4A2C")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elems.append(totals)
    elems.append(Spacer(1, 10*mm))

    if order.get("notes"):
        elems.append(Paragraph(f"<b>Notes:</b> {order['notes']}", sub))
        elems.append(Spacer(1, 4*mm))

    elems.append(Paragraph(
        "Thank you for choosing Green Groceries. Payment due within 7 days.",
        ParagraphStyle("foot", parent=styles["Normal"], fontSize=9,
                       textColor=colors.HexColor("#747A76"), alignment=1)))

    doc_pdf.build(elems)
    out.seek(0)
    filename = f"invoice_{inv_no}.pdf"
    return StreamingResponse(out, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _pdf_purchase_sheet(sheet: dict):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from io import BytesIO

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Heading1"], textColor=colors.HexColor("#1F4A2C"))
    sub = ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.HexColor("#747A76"))

    elems = [
        Paragraph("Green Groceries — Combined Purchase Sheet", h),
        Paragraph(
            f"Period: <b>{sheet['date_from']}</b> to <b>{sheet['date_to']}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Hotels: <b>{sheet['total_hotels']}</b> &nbsp;|&nbsp; Orders: <b>{sheet['total_orders']}</b>", sub),
        Spacer(1, 8*mm),
    ]
    data = [["Item", "Category", "Unit", "Total Qty", "Hotels"]]
    for r in sheet["rows"]:
        data.append([r["name"], r["category"].title(), r["unit"], str(r["total_quantity"]), str(r["hotel_count"])])
    if len(data) == 1:
        data.append(["—", "—", "—", "—", "—"])
    t = Table(data, repeatRows=1, hAlign="LEFT", colWidths=[55*mm, 30*mm, 20*mm, 30*mm, 25*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4A2C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (3, 1), (4, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F7F5F0")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E0D8")),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    elems.append(t)
    doc.build(elems)
    out.seek(0)
    filename = f"purchase_sheet_{sheet['date_from']}_to_{sheet['date_to']}.pdf"
    return StreamingResponse(out, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _pdf_orders(docs: List[dict], label: str, date_from, date_to):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from io import BytesIO

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Heading1"], textColor=colors.HexColor("#1F4A2C"))
    sub = ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.HexColor("#747A76"))

    elems = [
        Paragraph("Green Groceries — Orders Report", h),
        Paragraph(
            f"Hotel: <b>{label}</b> &nbsp;|&nbsp; From: <b>{date_from or '—'}</b> &nbsp;|&nbsp; To: <b>{date_to or '—'}</b>",
            sub),
        Spacer(1, 6*mm),
    ]

    data = [["Date", "Hotel", "Item", "Unit", "Qty", "Status"]]
    for d in docs:
        for ln in d["lines"]:
            data.append([d["order_date"], d["hotel_name"], ln["name"], ln["unit"],
                         str(ln["quantity"]), d["status"]])
    if len(data) == 1:
        data.append(["—"] * 6)
    t = Table(data, repeatRows=1, hAlign="LEFT",
              colWidths=[25*mm, 40*mm, 40*mm, 18*mm, 18*mm, 25*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4A2C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (4, 1), (4, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F7F5F0")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E0D8")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elems.append(t)
    doc.build(elems)
    out.seek(0)
    filename = f"orders_{label}.pdf"
    return StreamingResponse(out, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# ----- Analytics (admin) -----
@api.get("/analytics")
async def analytics(_: dict = Depends(require_admin), days: int = 7):
    today = datetime.now(timezone.utc).date()
    start = today - timedelta(days=days - 1)
    start_str = start.strftime("%Y-%m-%d")
    today_str = today.strftime("%Y-%m-%d")

    docs = await db.orders.find(
        {"order_date": {"$gte": start_str, "$lte": today_str}, "status": {"$ne": "cancelled"}}
    ).to_list(5000)

    # Build daily volume series
    by_day: dict = {}
    for i in range(days):
        d = (start + timedelta(days=i)).strftime("%Y-%m-%d")
        by_day[d] = {"date": d, "orders": 0, "quantity": 0.0}
    item_totals: dict = {}
    hotel_totals: dict = {}
    for d in docs:
        key = d["order_date"]
        if key in by_day:
            by_day[key]["orders"] += 1
            for ln in d["lines"]:
                by_day[key]["quantity"] += float(ln["quantity"])
        for ln in d["lines"]:
            n = ln["name"]
            if n not in item_totals:
                item_totals[n] = {"name": n, "unit": ln["unit"], "category": ln["category"], "quantity": 0.0}
            item_totals[n]["quantity"] += float(ln["quantity"])
        hid = d["hotel_id"]
        if hid not in hotel_totals:
            hotel_totals[hid] = {"hotel_id": hid, "hotel_name": d["hotel_name"], "orders": 0, "quantity": 0.0}
        hotel_totals[hid]["orders"] += 1
        for ln in d["lines"]:
            hotel_totals[hid]["quantity"] += float(ln["quantity"])

    weekly = sorted(by_day.values(), key=lambda x: x["date"])
    for w in weekly:
        w["quantity"] = round(w["quantity"], 2)
    top_items = sorted(item_totals.values(), key=lambda x: x["quantity"], reverse=True)[:8]
    for it in top_items:
        it["quantity"] = round(it["quantity"], 2)
    top_hotels = sorted(hotel_totals.values(), key=lambda x: x["quantity"], reverse=True)[:8]
    for h in top_hotels:
        h["quantity"] = round(h["quantity"], 2)

    return {
        "range": {"from": start_str, "to": today_str, "days": days},
        "weekly_volume": weekly,
        "top_items": top_items,
        "top_hotels": top_hotels,
    }


# ----- Stats -----
@api.get("/stats")
async def stats(_: dict = Depends(require_admin)):
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    total_hotels = await db.users.count_documents({"role": "hotel"})
    pending_orders = await db.orders.count_documents({"status": "pending"})
    today_orders_docs = await db.orders.find({"order_date": today}).to_list(2000)
    total_qty_today = 0.0
    for d in today_orders_docs:
        for ln in d["lines"]:
            total_qty_today += float(ln["quantity"])
    total_orders = await db.orders.count_documents({})
    return {
        "total_hotels": total_hotels,
        "pending_orders": pending_orders,
        "total_orders": total_orders,
        "today_total_quantity": round(total_qty_today, 2),
        "today_orders_count": len(today_orders_docs),
    }


# ----- Hotels (admin only) -----
class HotelCreateIn(BaseModel):
    email: EmailStr
    password: str = Field(min_length=6)
    name: str = Field(min_length=2)
    hotel_name: str = Field(min_length=2)
    phone: Optional[str] = None
    address: Optional[str] = None


class HotelPasswordIn(BaseModel):
    password: str = Field(min_length=6)


@api.post("/hotels")
async def create_hotel(payload: HotelCreateIn, _: dict = Depends(require_admin)):
    email = payload.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email already registered")
    doc = {
        "email": email,
        "password_hash": hash_password(payload.password),
        "name": payload.name,
        "role": "hotel",
        "hotel_name": payload.hotel_name,
        "phone": payload.phone,
        "address": payload.address,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    result = await db.users.insert_one(doc)
    return {
        "id": str(result.inserted_id),
        "email": email,
        "name": payload.name,
        "hotel_name": payload.hotel_name,
        "phone": payload.phone,
        "address": payload.address,
    }


@api.patch("/hotels/{hotel_id}/password")
async def reset_hotel_password(hotel_id: str, payload: HotelPasswordIn, _: dict = Depends(require_admin)):
    hotel = await db.users.find_one({"_id": ObjectId(hotel_id), "role": "hotel"})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel not found")
    await db.users.update_one(
        {"_id": ObjectId(hotel_id)},
        {"$set": {"password_hash": hash_password(payload.password)}},
    )
    return {"ok": True}


@api.delete("/hotels/{hotel_id}")
async def delete_hotel(hotel_id: str, _: dict = Depends(require_admin)):
    result = await db.users.delete_one({"_id": ObjectId(hotel_id), "role": "hotel"})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Hotel not found")
    return {"ok": True}


@api.get("/hotels")
async def list_hotels(_: dict = Depends(require_admin)):
    docs = await db.users.find({"role": "hotel"}).sort("hotel_name", 1).to_list(1000)
    return [{
        "id": str(d["_id"]), "name": d["name"], "email": d["email"],
        "hotel_name": d.get("hotel_name"), "phone": d.get("phone"),
        "address": d.get("address"),
    } for d in docs]


app.include_router(api)

origins_env = os.environ.get("CORS_ORIGINS", "*")
allow_origins = ["*"] if origins_env == "*" else origins_env.split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=allow_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)


@app.on_event("startup")
async def on_startup():
    await db.users.create_index("email", unique=True)
    await db.items.create_index("name", unique=True)
    await db.orders.create_index([("order_date", 1)])
    await db.orders.create_index([("hotel_id", 1)])

    # Seed admin
    admin_email = os.environ["ADMIN_EMAIL"].lower()
    admin_password = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "email": admin_email,
            "password_hash": hash_password(admin_password),
            "name": "Admin",
            "role": "admin",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info("Seeded admin user: %s", admin_email)
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one(
            {"email": admin_email},
            {"$set": {"password_hash": hash_password(admin_password)}}
        )
        logger.info("Updated admin password hash")

    # Seed default catalog
    if await db.items.count_documents({}) == 0:
        default_items = [
            {"name": "Potato", "category": "vegetable", "unit": "kg", "icon": "🥔"},
            {"name": "Onion", "category": "vegetable", "unit": "kg", "icon": "🧅"},
            {"name": "Tomato", "category": "vegetable", "unit": "kg", "icon": "🍅"},
            {"name": "Garlic", "category": "vegetable", "unit": "kg", "icon": "🧄"},
            {"name": "Carrot", "category": "vegetable", "unit": "kg", "icon": "🥕"},
            {"name": "Capsicum", "category": "vegetable", "unit": "kg", "icon": "🫑"},
            {"name": "Apple", "category": "fruit", "unit": "kg", "icon": "🍎"},
            {"name": "Banana", "category": "fruit", "unit": "dozen", "icon": "🍌"},
            {"name": "Orange", "category": "fruit", "unit": "kg", "icon": "🍊"},
            {"name": "Papaya", "category": "fruit", "unit": "kg", "icon": "🍈"},
        ]
        for it in default_items:
            it["created_at"] = datetime.now(timezone.utc).isoformat()
        await db.items.insert_many(default_items)
        logger.info("Seeded default catalog items")


@app.on_event("shutdown")
async def on_shutdown():
    client.close()
